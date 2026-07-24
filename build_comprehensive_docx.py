import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

doc = Document()

# Set Standard Page Margins (1 inch)
for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

# Helper Functions
def add_title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = RGBColor(16, 44, 87) # Dark Navy
    p.paragraph_format.space_after = Pt(6)
    return p

def add_subtitle(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(13)
    run.font.italic = True
    run.font.color.rgb = RGBColor(100, 116, 139) # Slate Gray
    p.paragraph_format.space_after = Pt(24)
    return p

def add_heading_1(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(15)
    run.font.bold = True
    run.font.color.rgb = RGBColor(16, 44, 87) # Dark Navy
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(8)
    return p

def add_heading_2(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(12.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor(30, 58, 138) # Royal Blue
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    return p

def add_paragraph(text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Arial"
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(30, 41, 59)
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(51, 65, 85)
    return p

def add_bullet(text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Arial"
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(30, 41, 59)
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(51, 65, 85)
    return p

def add_callout(text, title="PENTING"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    
    shading_xml = parse_xml(r'<w:shd {} w:fill="F1F5F9"/>'.format(nsdecls('w')))
    cell._tc.get_or_add_tcPr().append(shading_xml)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(r'''
        <w:tcBorders {} >
            <w:left w:val="single" w:sz="24" w:space="0" w:color="2563EB"/>
            <w:top w:val="none"/>
            <w:right w:val="none"/>
            <w:bottom w:val="none"/>
        </w:tcBorders>
    '''.format(nsdecls('w')))
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Inches(0.1)
    
    r_title = p.add_run(f"📌 {title}: ")
    r_title.font.name = "Arial"
    r_title.font.size = Pt(9.5)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(37, 99, 235)
    
    r_text = p.add_run(text)
    r_text.font.name = "Arial"
    r_text.font.size = Pt(9.5)
    r_text.font.color.rgb = RGBColor(30, 41, 59)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_image_with_caption(img_path, caption):
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(4)
        run = p_img.add_run()
        run.add_picture(img_path, width=Inches(6.2))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(12)
        r_cap = p_cap.add_run(f"Gambar: {caption}")
        r_cap.font.name = "Arial"
        r_cap.font.size = Pt(9)
        r_cap.font.italic = True
        r_cap.font.color.rgb = RGBColor(100, 116, 139)

print("Starting comprehensive document generation...")

# Document Header / Title
add_title("PANDUAN LENGKAP & OPERASIONAL SISTEM AUDIT & INSPEKSI GMP")
add_subtitle("Dokumentasi Terperinci Fitur, Antarmuka, dan Petunjuk Penggunaan untuk Administrator, Auditor, dan Penanggung Jawab Area (Auditee)")

doc.add_paragraph().paragraph_format.space_after = Pt(6)

# =========================================================================
# BAB 1: GAMBARAN UMUM
# =========================================================================
add_heading_1("BAB 1: GAMBARAN UMUM & STRUKTUR PERAN SISTEM")
add_paragraph("Sistem Audit & Inspeksi GMP (Good Manufacturing Practice) ini dikembangkan sebagai platform digital terpusat untuk mengelola seluruh siklus pengawasan mutu, kebersihan, dan keselamatan kerja fasilitas pabrik PT Dairy Tbk.")

add_paragraph("Aplikasi ini menghubungkan tiga peran utama dengan alur kerja yang tertata rapi:")
add_bullet("Memiliki akses penuh untuk mengelola struktur data induk (Area, Kawasan, Aspek, Uraian), akun pengguna, hak akses, konfigurasi batas sistem, kustomisasi email, ekspor laporan Excel QC, serta memantau audit trail aktivitas.", "1. Peran Administrator Sistem: ")
add_bullet("Bertanggung jawab melakukan inspeksi fisik di lapangan menggunakan formulir digital instan, menilai kepatuhan (Nilai 0 atau 2), mengunggah bukti foto ketidaksesuaian, dan menyelesaikan sesi inspeksi.", "2. Peran Auditor Quality Assurance (QA): ")
add_bullet("Bertanggung jawab menerima notifikasi temuan (issue), melaksanakan tindakan perbaikan di lokasi, dan mengunggah foto bukti perbaikan sebelum tenggat waktu berakhir.", "3. Peran Penanggung Jawab Area (Auditee / PIC): ")

add_callout("Setiap aksi perubahan data dan status pemeriksaan dicatat secara permanen di audit trail backend, serta dilindungi oleh otentikasi JWT dan middleware kontrol akses (RBAC).", "PRINSIP KEAMANAN & TRANSPARANSI")

# =========================================================================
# BAB 2: MODUL LOGIN
# =========================================================================
add_heading_1("BAB 2: MODUL MASUK SISTEM (LOGIN & OTENTIKASI)")
add_paragraph("Sebelum dapat mengakses fitur aplikasi, pengguna diwajibkan melakukan autentikasi login melalui formulir otentikasi terenkripsi.")

add_image_with_caption("/Users/apple/System-Audit-/screenshots/01_login_page.png", "1. Halaman Utama Autentikasi Pengguna (Login)")

add_heading_2("Langkah-Langkah Masuk Sistem:")
add_bullet("Akses URL portal resmi perusahaan melalui peramban web.", "• Langkah 1: ")
add_bullet("Masukkan Nama Pengguna (Username) resmi Anda (contoh: admin, auditor, atau amcu).", "• Langkah 2: ")
add_bullet("Masukkan Kata Sandi (Password) akun Anda.", "• Langkah 3: ")
add_bullet("Klik tombol 'Masuk Ke Sistem'. Sistem akan memverifikasi kredensial Anda dan mengarahkan Anda ke Dasbor Utama yang sesuai dengan Peran Anda.", "• Langkah 4: ")

# =========================================================================
# BAB 3: MODUL ADMINISTRATOR (SEMUA HALAMAN ADMIN)
# =========================================================================
add_heading_1("BAB 3: MODUL ADMINISTRATOR (DASBOR & KELOLA FITUR)")
add_paragraph("Administrator memiliki wewenang pengawasan tertinggi. Berikut adalah dokumentasi lengkap seluruh halaman dan fitur yang tersedia untuk peran Administrator:")

# 3.1 Dasbor Admin
add_heading_2("3.1 Dasbor Administrator - Kartu Ringkasan Performa")
add_paragraph("Halaman utama Dasbor Admin menyajikan kartu ringkasan indikator performa utama (KPI) secara real-time.")
add_image_with_caption("/Users/apple/System-Audit-/screenshots/02_admin_dashboard_top.png", "2. Dasbor Administrator - Kartu Ringkasan Indikator Utama")

add_bullet("Menampilkan jumlah pemeriksaan yang sedang aktif dijalankan oleh Auditor.", "• Inspeksi Berlangsung: ")
add_bullet("Menampilkan persentase kepatuhan mutu rata-rata dari seluruh area pabrik.", "• Tingkat Kepatuhan: ")
add_bullet("Jumlah ketidaksesuaian yang membutuhkan tindakan perbaikan dari PIC Area.", "• Temuan Terbuka: ")
add_bullet("Jumlah temuan perbaikan yang telah melampaui batas tenggat waktu penyelesaian.", "• Temuan Jatuh Tempo: ")

add_heading_2("3.2 Dasbor Administrator - Pengawasan per Peran (Role Progress)")
add_paragraph("Bagian pengawasan peran memantau efisiensi pelaksanaan audit oleh Auditor dan progres perbaikan temuan oleh PIC Area.")
add_image_with_caption("/Users/apple/System-Audit-/screenshots/03_admin_dashboard_monitoring.png", "3. Dasbor Administrator - Pengawasan Aktivitas Auditor & PIC Area")

add_heading_2("3.3 Dasbor Administrator - Grafik Tren Kepatuhan Multi-Periode")
add_paragraph("Grafik tren kepatuhan memungkinkan Administrator menganalisis pergerakan mutu dalam rentang waktu yang dapat disesuaikan (1 Bulan, 3 Bulan, 6 Bulan, atau 1 Tahun).")
add_image_with_caption("/Users/apple/System-Audit-/screenshots/04_admin_dashboard_chart.png", "4. Dasbor Administrator - Grafik Tren Tingkat Kepatuhan Multi-Periode")

# 3.4 Form & Data Inspeksi
add_heading_2("3.4 Halaman Form & Data Inspeksi Digital")
add_paragraph("Halaman ini mencatat seluruh dokumen inspeksi yang dibuat. Administrator dapat memantau status inspeksi (Draft, Ongoing, Completed, Approved) dan mengunduh laporan.")
add_image_with_caption("/Users/apple/System-Audit-/screenshots/05_admin_inspections_list.png", "5. Halaman Daftar Form & Status Inspeksi Digital")

# 3.5 Manajemen Temuan & WO/WR
add_heading_2("3.5 Halaman Manajemen Temuan (Issues)")
add_paragraph("Seluruh temuan ketidaksesuaian hasil pemeriksaan terdaftar di halaman ini. Administrator dapat memantau progres tindak lanjut PIC, meninjau foto bukti perbaikan, dan memberikan persetujuan.")
add_image_with_caption("/Users/apple/System-Audit-/screenshots/06_admin_issues_list.png", "6. Halaman Daftar Temuan Inspeksi (Issues)")

add_heading_2("3.6 Halaman Perintah Kerja (Work Order / Work Request)")
add_paragraph("Apabila temuan membutuhkan perbaikan fisik skala besar atau keterlibatan tim mekanik/teknisi, Administrator dapat memantau status pengajuan WO/WR di halaman ini.")
add_image_with_caption("/Users/apple/System-Audit-/screenshots/07_admin_wowr_page.png", "7. Halaman Kelola Perintah Kerja (WO / WR)")

# 3.7 Data Induk (Master Data)
add_heading_2("3.7 Kelola Data Induk - Tab Area & Kawasan")
add_paragraph("Mengatur hierarki lokasi fisik pabrik mulai dari Area Utama hingga Kawasan dan Detail Kawasan Spesifik.")
add_image_with_caption("/Users/apple/System-Audit-/screenshots/08_admin_master_area.png", "8. Data Induk - Pengaturan Area & Kawasan Pabrik")

add_heading_2("3.8 Kelola Data Induk - Tab Aspek Inspeksi")
add_paragraph("Mengelompokkan standar pemeriksaan ke dalam Kategori Aspek Mutu (seperti Kebersihan, Keamanan Pangan, Fasilitas, Bangunan).")
add_image_with_caption("/Users/apple/System-Audit-/screenshots/09_admin_master_aspek.png", "9. Data Induk - Kelola Kategori Aspek Inspeksi")

add_heading_2("3.9 Kelola Data Induk - Tab Uraian Pertanyaan Checklist")
add_paragraph("Daftar rinci pertanyaan checklist yang harus diperiksa oleh Auditor di lapangan.")
add_image_with_caption("/Users/apple/System-Audit-/screenshots/10_admin_master_uraian.png", "10. Data Induk - Kelola Uraian Pertanyaan Checklist")

# 3.10 Users Management
add_heading_2("3.10 Halaman Manajemen Pengguna (Users)")
add_paragraph("Administrator dapat mendaftarkan akun pengguna baru, menentukan kata sandi, dan mengatur Hak Akses / Role (Admin, Auditor, Auditee/PIC).")
add_image_with_caption("/Users/apple/System-Audit-/screenshots/11_admin_users_list.png", "11. Halaman Manajemen Pengguna & Hak Akses")

# 3.11 GMP Data
add_heading_2("3.11 Halaman Rekapitulasi Data Inspeksi GMP & Ekspor Excel")
add_paragraph("Halaman ini menyajikan rekapitulasi data hasil audit GMP secara komprehensif dan menyediakan tombol ekspor laporan resmi berbasis template Excel (.xlsx).")
add_image_with_caption("/Users/apple/System-Audit-/screenshots/12_admin_gmp_data.png", "12. Halaman Data Inspeksi GMP & Ekspor Laporan Excel")

# 3.12 Audit Trail
add_heading_2("3.12 Audit Trail - Tab Riwayat Aktivitas API")
add_paragraph("Mencatat secara rinci setiap aksi penambahan, pengeditan, atau penghapusan data di dalam sistem.")
add_image_with_caption("/Users/apple/System-Audit-/screenshots/13_admin_logs_activity.png", "13. Audit Trail - Riwayat Aktivitas & Permintaan API")

add_heading_2("3.13 Audit Trail - Tab Riwayat Sesi Masuk Pengguna")
add_paragraph("Memantau log waktu login, logout, alamat IP, dan perangkat yang digunakan pengguna.")
add_image_with_caption("/Users/apple/System-Audit-/screenshots/14_admin_logs_login.png", "14. Audit Trail - Riwayat Sesi Masuk Pengguna")

# 3.14 Settings
add_heading_2("3.14 Pengaturan Sistem - Tab Pengaturan Umum")
add_paragraph("Konfigurasi parameter umum seperti batas percobaan login, waktu timeout sesi, dan ukuran maksimal unggahan file.")
add_image_with_caption("/Users/apple/System-Audit-/screenshots/15_admin_settings_general.png", "15. Pengaturan Sistem - Konfigurasi Parameter Umum")

add_heading_2("3.15 Pengaturan Sistem - Tab Template Email & Visual Editor")
add_paragraph("Memungkinkan Administrator menyesuaikan isi dan tampilan surat pemberitahuan email otomatis menggunakan Visual Email Editor interaktif.")
add_image_with_caption("/Users/apple/System-Audit-/screenshots/16_admin_settings_email.png", "16. Pengaturan Sistem - Template Email & Editor Visual")

# =========================================================================
# BAB 4: MODUL AUDITOR
# =========================================================================
add_heading_1("BAB 4: MODUL & PANDUAN PENGGUNAAN AUDITOR")
add_paragraph("Auditor memiliki antarmuka khusus yang dirancang untuk mempermudah pemeriksaan fisik di lapangan.")

add_heading_2("4.1 Dasbor Utama Auditor")
add_paragraph("Menampilkan jumlah inspeksi milik auditor yang sedang berlangsung, inspeksi yang telah selesai, serta jumlah temuan terbuka.")
add_image_with_caption("/Users/apple/System-Audit-/screenshots/17_auditor_dashboard_overview.png", "17. Dasbor Utama Peran Auditor")

add_heading_2("4.2 Grafik Tren Inspeksi Bulanan Auditor")
add_paragraph("Grafik frekuensi pemeriksaan yang diselesaikan oleh Auditor setiap bulannya.")
add_image_with_caption("/Users/apple/System-Audit-/screenshots/18_auditor_dashboard_trend.png", "18. Grafik Tren Inspeksi Bulanan Auditor")

add_heading_2("Alur Kerja Pemeriksaan oleh Auditor:")
add_bullet("Auditor memilih lokasi Kawasan dan Detail Kawasan yang akan diperiksa.", "1. Pilih Lokasi: ")
add_bullet("Memeriksa setiap item uraian. Jika sesuai berikan Nilai 2, jika ditemukan ketidaksesuaian berikan Nilai 0.", "2. Penilaian Checklist: ")
add_bullet("Apabila Nilai 0 dipilih, Auditor wajib mengunggah foto bukti temuan lapangan dan mengisi kolom keterangan.", "3. Unggah Bukti Foto: ")
add_bullet("Setelah selesai, Auditor menekan 'Selesai Inspeksi'. Jika seluruh kawasan dalam area selesai, sistem otomatis mengirim email laporan ke PIC Area.", "4. Selesaikan Inspeksi: ")

# =========================================================================
# BAB 5: MODUL AUDITEE / PIC AREA
# =========================================================================
add_heading_1("BAB 5: MODUL & PANDUAN PENGGUNAAN AUDITEE (PIC AREA)")
add_paragraph("Penanggung Jawab Area (PIC) difasilitasi dengan dasbor yang berfokus pada daftar tugas tindakan perbaikan temuan.")

add_heading_2("5.1 Dasbor Utama Penanggung Jawab Area")
add_paragraph("Menampilkan ringkasan total tugas temuan, tugas terbuka, tugas menunggu validasi, dan tugas yang telah selesai disetujui.")
add_image_with_caption("/Users/apple/System-Audit-/screenshots/19_auditee_dashboard_overview.png", "19. Dasbor Utama Peran Auditee / PIC Area")

add_heading_2("5.2 Daftar Temuan Prioritas & Tenggat Waktu")
add_paragraph("Daftar urutan temuan yang harus diperbaiki PIC berdasarkan kedekatan batas tenggat waktu.")
add_image_with_caption("/Users/apple/System-Audit-/screenshots/20_auditee_tasks_list.png", "20. Daftar Temuan Prioritas & Tenggat Waktu Perbaikan")

add_heading_2("Alur Tindak Lanjut Perbaikan oleh PIC:")
add_bullet("PIC menerima email pemberitahuan otomatis saat area miliknya selesai diinspeksi dan ditemukan masalah.", "1. Notifikasi Email: ")
add_bullet("PIC masuk ke dasbor dan meninjau daftar temuan terbuka.", "2. Tinjau Tugas: ")
add_bullet("Setelah perbaikan fisik dilakukan di lokasi, PIC mengunggah foto bukti perbaikan dan memberikan deskripsi tindakan penanganan.", "3. Unggah Perbaikan: ")
add_bullet("Status berubah menjadi 'Pending Validation' untuk diperiksa kembali oleh Auditor/Admin hingga dinyatakan 'Closed'.", "4. Verifikasi Penutupan: ")

# =========================================================================
# BAB 6: PENUTUP
# =========================================================================
add_heading_1("BAB 6: PENUTUP")
add_paragraph("Dengan sistem yang saling terintegrasi ini, PT Dairy Tbk dapat mempertahankan standar kepatuhan mutu GMP secara konsisten, cepat, dan transparan. Apabila terdapat kendala teknis, silakan hubungi tim Administrator Sistem.")

# Save Document
output_file = "/Users/apple/System-Audit-/Panduan_Penggunaan_Sistem_Audit_Cimory.docx"
doc.save(output_file)
print(f"Comprehensive document created successfully at: {output_file}")
