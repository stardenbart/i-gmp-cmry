import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

doc = Document()

# Set Standard Page Margins (1 inch)
sections = doc.sections
for section in sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

# Style Helper Functions
def add_title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = RGBColor(16, 44, 87) # Dark Navy
    p.paragraph_format.space_after = Pt(6)
    return p

def add_subtitle(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(13)
    run.font.italic = True
    run.font.color.rgb = RGBColor(100, 116, 139) # Slate Gray
    p.paragraph_format.space_after = Pt(24)
    return p

def add_heading_1(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = RGBColor(16, 44, 87) # Dark Navy
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(8)
    return p

def add_heading_2(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = RGBColor(30, 58, 138) # Royal Blue
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    return p

def add_paragraph(text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Arial"
        r_pre.font.size = Pt(10.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(30, 41, 59)
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor(51, 65, 85)
    return p

def add_bullet(text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Arial"
        r_pre.font.size = Pt(10.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(30, 41, 59)
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor(51, 65, 85)
    return p

def add_callout(text, title="PENTING"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    
    # Border & Fill styling
    shading_xml = parse_xml(r'<w:shd {} w:fill="F1F5F9"/>'.format(nsdecls('w')))
    cell._tc.get_or_add_tcPr().append(shading_xml)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(r'''
        <w:tcBorders {} >
            <w:left w:val="single" w:sz="24" w:space="0" w:color="2563EB"/>
            <w:top w:val="none"/>
            <w:right w:val="none"/>
            <w:bottom w:val="none"/>
        </w:tcBorders>
    '''.format(nsdecls('w')))
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Inches(0.1)
    
    r_title = p.add_run(f"📌 {title}: ")
    r_title.font.name = "Arial"
    r_title.font.size = Pt(10)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(37, 99, 235)
    
    r_text = p.add_run(text)
    r_text.font.name = "Arial"
    r_text.font.size = Pt(10)
    r_text.font.color.rgb = RGBColor(30, 41, 59)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_image_with_caption(img_path, caption):
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(10)
        p_img.paragraph_format.space_after = Pt(4)
        run = p_img.add_run()
        run.add_picture(img_path, width=Inches(6.2))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(14)
        r_cap = p_cap.add_run(f"Gambar: {caption}")
        r_cap.font.name = "Arial"
        r_cap.font.size = Pt(9.5)
        r_cap.font.italic = True
        r_cap.font.color.rgb = RGBColor(100, 116, 139)

print("Starting document creation...")

# Document Cover / Header
add_title("PANDUAN OPERASIONAL & PENGGUNAAN SISTEM AUDIT & INSPEKSI GMP")
add_subtitle("Dokumentasi Sistem Pengawasan Audit Kepatuhan PT Dairy Tbk - Edisi 2026")

# Line separator
p_hr = doc.add_paragraph()
p_hr.paragraph_format.space_after = Pt(12)

# Bab 1
add_heading_1("BAB 1: GAMBARAN UMUM & TUJUAN SISTEM")
add_paragraph("Sistem Audit & Inspeksi GMP (Good Manufacturing Practice) ini dirancang khusus untuk memfasilitasi proses pemeriksaan kepatuhan standar mutu dan keselamatan fasilitas pabrik. Sistem ini menghubungkan tiga peran utama dalam satu platform digital terintegrasi:")

add_bullet("Memiliki akses penuh untuk mengelola pengguna, data induk lokasi/aspek inspeksi, mengekspor laporan Excel (.xlsx), mengonfigurasi template email, serta memantau seluruh aktivitas sistem melalui riwayat audit.", "1. Administrator Sistem: ")
add_bullet("Bertanggung jawab melakukan verifikasi dan pemeriksaan fisik secara langsung di lapangan menggunakan formulir digital instan, serta mengunggah bukti gambar jika ditemukan ketidaksesuaian.", "2. Auditor Quality Assurance (QA): ")
add_bullet("Bertanggung jawab menerima notifikasi penugasan perbaikan (issue), melakukan tindak lanjut di area produksi/fasilitas, dan mengunggah bukti perbaikan sebelum tenggat waktu berakhir.", "3. Penanggung Jawab Area (Auditee / PIC): ")

add_callout("Seluruh aktivitas transaksi data, perubahan status inspeksi, dan sesi masuk pengguna dicatat secara permanen di dalam sistem untuk menjamin transparansi dan akuntabilitas.", "PRINSIP UTAMA")

# Bab 2
add_heading_1("BAB 2: MODUL MASUK SISTEM (LOGIN & KEAMANAN)")
add_paragraph("Untuk menjaga kerahasiaan data perusahaan, seluruh pengguna diwajibkan melakukan autentikasi sebelum dapat mengakses dasbor dan fitur aplikasi.")

add_image_with_caption("/Users/apple/System-Audit-/screenshots/01_login_page.png", "Halaman Autentikasi Pengguna (Login)")

add_heading_2("Langkah-Langkah Masuk Sistem:")
add_bullet("Buka peramban (browser) dan akses alamat URL portal resmi perusahaan.", "1. Akses URL: ")
add_bullet("Masukkan alamat surat elektronik (email) yang telah terdaftar resmi oleh Administrator.", "2. Input Email: ")
add_bullet("Masukkan kata sandi akun Anda.", "3. Input Password: ")
add_bullet("Klik tombol 'Masuk Ke Sistem'. Sistem akan secara otomatis mengarahkan Anda ke Dasbor Utama sesuai dengan peran (Role) akun Anda.", "4. Klik Masuk: ")

# Bab 3
add_heading_1("BAB 3: DASBOR UTAMA ADMINISTRATOR (ADMIN DASHBOARD)")
add_paragraph("Dasbor Administrator memberikan ringkasan pengawasan tingkat tinggi mengenai seluruh kegiatan inspeksi dan tindak lanjut temuan di pabrik secara real-time.")

add_image_with_caption("/Users/apple/System-Audit-/screenshots/02_admin_dashboard.png", "Tampilan Utama Dasbor Administrator")

add_heading_2("Komponen Utama Dasbor Administrator:")
add_bullet("Menampilkan jumlah pemeriksaan yang sedang aktif dilakukan oleh para auditor di lapangan.", "• Kartu Inspeksi Berlangsung: ")
add_bullet("Menampilkan persentase rata-rata kepatuhan standar mutu secara keseluruhan.", "• Kartu Tingkat Kepatuhan (Compliance Rate): ")
add_bullet("Jumlah temuan ketidaksesuaian yang membutuhkan tindakan perbaikan dari PIC.", "• Kartu Temuan Terbuka: ")
add_bullet("Jumlah temuan yang telah melewati batas tenggat waktu penyelesaian.", "• Kartu Temuan Jatuh Tempo (Overdue): ")
add_bullet("Grafik interaktif yang menampilkan dinamika pergerakan kepatuhan. Pengguna dapat memilih rentang waktu 1 Bulan, 3 Bulan, 6 Bulan, atau 1 Tahun.", "• Grafik Tren Tingkat Kepatuhan: ")
add_bullet("Mengawasi rasio penyelesaian inspeksi oleh Auditor serta persentase penanganan temuan oleh PIC Area.", "• Pengawasan per Peran (Role Progress): ")

# Bab 4
add_heading_1("BAB 4: FITUR KELOLA & MANAJEMEN SISTEM")

add_heading_2("4.1 Halaman Form & Data Inspeksi")
add_paragraph("Halaman ini menampilkan seluruh daftar dokumen inspeksi yang tercatat di dalam sistem beserta statusnya (Draft, Berlangsung, Selesai, atau Disetujui).")
add_image_with_caption("/Users/apple/System-Audit-/screenshots/03_inspections_page.png", "Daftar Form & Status Inspeksi Digital")

add_heading_2("4.2 Manajemen Temuan Inspeksi (Issue & WO/WR)")
add_paragraph("Halaman Temuan mencatat setiap ketidaksesuaian yang ditemukan saat pemeriksaan. Dari halaman ini, Administrator dan Auditor dapat mengawasi status tindak lanjut, tenggat waktu, dan pembuatan Perintah Kerja (WO/WR).")
add_image_with_caption("/Users/apple/System-Audit-/screenshots/04_issues_page.png", "Daftar Temuan Inspeksi & Tindak Lanjut")

add_heading_2("4.3 Kelola Data Induk (Master Data)")
add_paragraph("Data Induk merupakan pondasi struktur pemeriksaan yang mencakup susunan Area, Kawasan, Detail Kawasan, Aspek Inspeksi, dan Uraian Pertanyaan.")
add_image_with_caption("/Users/apple/System-Audit-/screenshots/05_master_data_page.png", "Kelola Data Induk (Area, Kawasan, Aspek, & Uraian)")

add_heading_2("4.4 Kelola Manajemen Pengguna (Users)")
add_paragraph("Administrator dapat menambah akun pengguna baru, mengedit informasi profil, serta menetapkan peran (Role) seperti Administrator, Auditor, atau Penanggung Jawab (PIC).")
add_image_with_caption("/Users/apple/System-Audit-/screenshots/06_users_page.png", "Kelola Pengguna & Hak Akses Sistem")

add_heading_2("4.5 Data Inspeksi GMP & Ekspor Excel (.xlsx)")
add_paragraph("Laporan rekapitulasi nilai kepatuhan GMP dapat diunduh secara resmi dalam bentuk format Excel (.xlsx) yang rapi dan terstruktur sesuai standar formulir QC.")
add_image_with_caption("/Users/apple/System-Audit-/screenshots/07_gmp_data_page.png", "Halaman Rekapitulasi Data Inspeksi GMP")

add_heading_2("4.6 Riwayat Aktivitas & Log Masuk (Audit Trail)")
add_paragraph("Seluruh riwayat aksi pengguna, perubahan data, serta riwayat waktu masuk/keluar dicatat secara transparan di halaman Riwayat Aktivitas.")
add_image_with_caption("/Users/apple/System-Audit-/screenshots/08_logs_page.png", "Halaman Riwayat Aktivitas & Sesi Masuk (Audit Logs)")

add_heading_2("4.7 Pengaturan Sistem & Kustomisasi Email")
add_paragraph("Fitur kustomisasi template email notifikasi yang dilengkapi dengan Visual Email Editor interaktif untuk mempermudah pengaturan tampilan surat pemberitahuan otomatis.")
add_image_with_caption("/Users/apple/System-Audit-/screenshots/09_settings_page.png", "Pengaturan Konfigurasi Umum & Template Email")

# Bab 5
add_heading_1("BAB 5: PANDUAN PENGGUNAAN UNTUK AUDITOR")
add_paragraph("Auditor memiliki antarmuka khusus yang terfokus pada efisiensi pelaksanaan pemeriksaan fisik di lapangan.")

add_image_with_caption("/Users/apple/System-Audit-/screenshots/10_auditor_dashboard.png", "Tampilan Dasbor Khusus Peran Auditor")

add_heading_2("Alur Kerja Pemeriksaan oleh Auditor:")
add_bullet("Auditor memilih lokasi Kawasan dan Detail Kawasan yang akan diperiksa.", "1. Memulai Inspeksi: ")
add_bullet("Memilih opsi nilai pada setiap Uraian pertanyaan (Nilai 2 = Memenuhi / Tidak ada temuan, Nilai 0 = Terdapat temuan).", "2. Pengisian Form Checklist: ")
add_bullet("Apabila ditemukan ketidaksesuaian (Nilai 0), Auditor wajib mengunggah bukti foto lapangan dan memberikan keterangannya.", "3. Unggah Bukti Foto Temuan: ")
add_bullet("Setelah seluruh aspek terisi, Auditor menekan tombol 'Selesai Inspeksi'. Jika seluruh kawasan dalam area telah selesai, sistem akan secara otomatis mengirimkan email laporan kepada Penanggung Jawab (PIC).", "4. Penyelesaian Inspeksi: ")

# Bab 6
add_heading_1("BAB 6: PANDUAN PENGGUNAAN UNTUK AUDITEE (PIC AREA)")
add_paragraph("Penanggung Jawab Area (PIC) difasilitasi dengan dasbor yang berfokus pada daftar tugas perbaikan temuan.")

add_image_with_caption("/Users/apple/System-Audit-/screenshots/11_auditee_dashboard.png", "Tampilan Dasbor Khusus Peran Auditee (PIC Area)")

add_heading_2("Alur Tindak Lanjut Temuan oleh PIC:")
add_bullet("Setiap terjadi temuan di area yang menjadi tanggung jawabnya, PIC akan menerima notifikasi email otomatis.", "1. Penerimaan Notifikasi: ")
add_bullet("PIC masuk ke dasbor dan melihat daftar temuan yang perlu ditindaklanjuti berdasarkan prioritas tenggat waktu.", "2. Peninjauan Daftar Tugas: ")
add_bullet("Setelah melakukan perbaikan fisik di lokasi, PIC mengunggah bukti foto perbaikan dan memberikan catatan penjelasan.", "3. Pengunggahan Bukti Perbaikan: ")
add_bullet("Perbaikan yang diunggah akan diverifikasi oleh Auditor / Sistem untuk diubah statusnya menjadi Selesai (Closed).", "4. Validasi & Penutupan: ")

# Bab 7
add_heading_1("BAB 7: PENUTUP")
add_paragraph("Dengan hadirnya Sistem Audit & Inspeksi GMP ini, proses pengawasan kualitas dan kepatuhan fasilitas dapat berjalan lebih cepat, terstruktur, serta akuntabel. Apabila Anda mengalami kendala teknis atau membutuhkan bantuan lebih lanjut, silakan hubungi Tim IT Administrator.")

# Save Document
output_file = "/Users/apple/System-Audit-/Panduan_Penggunaan_Sistem_Audit_Cimory.docx"
doc.save(output_file)
print(f"Document created successfully at: {output_file}")
